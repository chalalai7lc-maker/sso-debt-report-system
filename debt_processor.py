import os
import re
import datetime
from typing import Optional, List, Dict, Any, Tuple
from collections import defaultdict
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from wcf_integrator import WCFIntegrator

THAI_MONTHS = [
    "", "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
    "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"
]

def format_period_ranges_b(period_list: List[str]) -> str:
    """
    แปลงรายการงวด 'MM/YYYY' ให้อยู่ในรูปแบบช่วงงวดภาษาไทย (แบบ B):
    - เดือนเดียว: เช่น 'กรกฎาคม 2569'
    - หลายเดือนต่อเนื่อง: เช่น 'มิถุนายน 2569 - กรกฎาคม 2569'
    - งวดไม่ต่อเนื่อง: แบ่งเป็นช่วงต่อเนื่องคั่นด้วยจุลภาค เช่น 'เมษายน 2569 - พฤษภาคม 2569, กรกฎาคม 2569'
    """
    parsed = []
    for p in period_list:
        if not p or not isinstance(p, str):
            continue
        parts = p.strip().split('/')
        if len(parts) == 2:
            try:
                m = int(parts[0])
                y = int(parts[1])
                if 1 <= m <= 12:
                    parsed.append((y, m))
            except ValueError:
                continue

    # เรียงลำดับตามปี พ.ศ. และเดือน และกำจัดค่าซ้ำ
    parsed = sorted(list(set(parsed)))
    if not parsed:
        return ""

    # จัดกลุ่มงวดที่ต่อเนื่องกัน
    groups: List[List[Tuple[int, int]]] = []
    current_group = [parsed[0]]

    for i in range(1, len(parsed)):
        prev_y, prev_m = current_group[-1]
        cur_y, cur_m = parsed[i]

        # ตรวจสอบว่าเดือนถัดไปต่อเนื่องกันหรือไม่
        month_diff = (cur_y - prev_y) * 12 + (cur_m - prev_m)
        if month_diff == 1:
            current_group.append((cur_y, cur_m))
        else:
            groups.append(current_group)
            current_group = [(cur_y, cur_m)]
    groups.append(current_group)

    # สร้างข้อความของแต่ละกลุ่ม
    formatted_groups = []
    for g in groups:
        if len(g) == 1:
            y, m = g[0]
            formatted_groups.append(f"{THAI_MONTHS[m]} {y}")
        else:
            start_y, start_m = g[0]
            end_y, end_m = g[-1]
            formatted_groups.append(f"{THAI_MONTHS[start_m]} {start_y} - {THAI_MONTHS[end_m]} {end_y}")

    return ", ".join(formatted_groups)


def parse_float_safe(val: Any) -> float:
    """แปลงค่าตัวเลขหรือข้อความจำนวนเงินที่มีเครื่องหมายจุลภาค (,) ให้เป็น float อย่างปลอดภัย"""
    if val is None or val == "":
        return 0.0
    if isinstance(val, (int, float)):
        return 0.0 if pd.isna(val) else float(val)
    cleaned = str(val).replace(",", "").strip()
    if not cleaned or cleaned in ("-", "None", "nan"):
        return 0.0
    try:
        return float(cleaned)
    except (ValueError, TypeError):
        return 0.0


class DebtReportProcessor:
    """
    คลาสหลักสำหรับการประมวลผลรายงานหนี้ค้างชำระของสถานประกอบการ (SSO Debt Report Processor)
    """

    def __init__(self, wcf_integrator: Optional[WCFIntegrator] = None):
        self.wcf_integrator = wcf_integrator or WCFIntegrator()

    def parse_debt_report(self, file_path: str) -> List[Dict[str, Any]]:
        """
        อ่านไฟล์รายงานหนี้ (.xlsx, .xls หรือ .csv) สกัดเอาเฉพาะแถวข้อมูลจริง
        ตัดส่วนหัวตาราง, หน้าซ้ำ (Page Break), และแถวสรุปยอดออก
        รองรับไฟล์ที่มีหัวกระดาษหลายบรรทัดและการจัดตำแหน่งคอลัมน์แบบไดนามิก
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"ไม่พบไฟล์: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()
        rows_matrix: List[List[Any]] = []

        if ext == ".xlsx":
            wb = openpyxl.load_workbook(file_path, data_only=True)
            ws = wb.active
            for r in range(1, ws.max_row + 1):
                rows_matrix.append([ws.cell(r, c).value for c in range(1, ws.max_column + 1)])
        elif ext == ".xls":
            df_raw = pd.read_excel(file_path, header=None, dtype=str).fillna("")
            rows_matrix = df_raw.values.tolist()
        elif ext == ".csv":
            df_raw = None
            for enc in ("utf-8-sig", "utf-8", "cp874", "tis-620"):
                try:
                    df_raw = pd.read_csv(file_path, header=None, dtype=str, encoding=enc, on_bad_lines="skip").fillna("")
                    break
                except Exception:
                    continue
            if df_raw is None:
                raise ValueError("ไม่สามารถอ่านไฟล์ CSV ได้ (Encoding ไม่รองรับ)")
            rows_matrix = df_raw.values.tolist()
        else:
            raise ValueError(f"ไม่รองรับนามสกุลไฟล์: {ext}")

        # ค้นหาบรรทัด Header ที่มีคำว่า 'เลขที่บัญชี'
        header_row_idx = 0
        col_map: Dict[str, int] = {}
        for r_idx, row in enumerate(rows_matrix[:20]):
            row_strs = [str(v or "").strip() for v in row]
            if any("เลขที่บัญชี" in s for s in row_strs):
                header_row_idx = r_idx
                for c_idx, val in enumerate(row_strs):
                    if "เลขที่บัญชี" in val and "acc" not in col_map:
                        col_map["acc"] = c_idx
                    elif val == "สาขา" or "เลขที่สาขา" in val:
                        col_map["branch"] = c_idx
                    elif "ชื่อสถานประกอบการ" in val:
                        col_map["name"] = c_idx
                    elif "สถานะ" in val and "สปก" in val:
                        col_map["status"] = c_idx
                    elif "งวดที่ค้าง" in val:
                        col_map["period"] = c_idx
                    elif "หนี้ประเมิน" in val:
                        col_map["est_debt"] = c_idx
                    elif "หนี้จริง" in val:
                        col_map["real_debt"] = c_idx
                    elif "เงินเพิ่ม" in val:
                        col_map["surcharge"] = c_idx
                    elif "รวมค้างชำระ" in val:
                        col_map["total_debt"] = c_idx
                    elif val == "ขั้นตอน" or "ขั้นตอนการติดตาม" in val:
                        col_map["step"] = c_idx
                    elif "วันที่บันทึก" in val:
                        col_map["step_date"] = c_idx
                    elif "ประเภทกิจการ" in val:
                        col_map["biz_type"] = c_idx
                break

        # ค่า Default ตำแหน่งคอลัมน์ตามโครงสร้างมาตรฐาน CN200090-1 (0-based index)
        c_acc = col_map.get("acc", 3)
        c_br = col_map.get("branch", 5)
        c_name = col_map.get("name", 6)
        c_status = col_map.get("status", 9)
        c_period = col_map.get("period", 12)
        c_est = col_map.get("est_debt", 13)
        c_real = col_map.get("real_debt", 14)
        c_sur = col_map.get("surcharge", 15)
        c_tot = col_map.get("total_debt", 16)
        c_step = col_map.get("step", 17)
        c_step_date = col_map.get("step_date", 18)
        c_biz = col_map.get("biz_type", 19)

        def _get_cell(row_list: List[Any], idx: int) -> Any:
            return row_list[idx] if 0 <= idx < len(row_list) else None

        records = []
        for r_idx in range(header_row_idx + 1, len(rows_matrix)):
            row = rows_matrix[r_idx]
            acc_raw = _get_cell(row, c_acc)
            if acc_raw is None:
                continue
            acc_str = str(acc_raw).strip()
            if acc_str.endswith(".0"):
                acc_str = acc_str[:-2]
            if not acc_str.isdigit() or acc_str in ("", "เลขที่บัญชี"):
                continue

            br_raw = _get_cell(row, c_br)
            branch_str = str(br_raw).strip() if br_raw is not None else ""
            if branch_str.endswith(".0"):
                branch_str = branch_str[:-2]
            branch_str = branch_str.zfill(6) if branch_str.isdigit() else branch_str

            name = str(_get_cell(row, c_name) or "").strip()
            status = str(_get_cell(row, c_status) or "").strip()
            period = str(_get_cell(row, c_period) or "").strip()

            est_debt = parse_float_safe(_get_cell(row, c_est))
            real_debt = parse_float_safe(_get_cell(row, c_real))
            surcharge = parse_float_safe(_get_cell(row, c_sur))
            raw_total = _get_cell(row, c_tot)
            total_debt = parse_float_safe(raw_total) if raw_total not in (None, "") else (est_debt + real_debt + surcharge)

            step = str(_get_cell(row, c_step) or "").strip()
            step_date = str(_get_cell(row, c_step_date) or "").strip()
            biz_type = str(_get_cell(row, c_biz) or "").strip()

            records.append({
                "row_index": r_idx + 1,
                "acc": acc_str,
                "branch": branch_str,
                "name_report": name,
                "status": status,
                "period": period,
                "est_debt": est_debt,
                "real_debt": real_debt,
                "surcharge": surcharge,
                "total_debt": total_debt,
                "step": step,
                "step_date": step_date,
                "biz_type": biz_type
            })

        return records

    def process(
        self,
        file_path: str,
        output_dir: Optional[str] = None,
        custom_filename: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        ประมวลผลข้อมูลรายงานหนี้:
        1. อ่านข้อมูล
        2. จัดกลุ่มตาม เลขที่บัญชี + สาขา
        3. เชื่อมโยงข้อมูล WCF Master Data
        4. จัดรูปแบบงวดเดือนค้างชำระ
        5. สรุปยอดหนี้ประเมิน หนี้จริง และยอดหนี้ไม่รวมเงินเพิ่ม
        6. ส่งออกเป็นไฟล์ Excel และคืนค่าผลลัพธ์เชิงสถิติ
        """
        raw_records = self.parse_debt_report(file_path)
        if not raw_records:
            return {
                "success": False,
                "error": "ไม่พบข้อมูลรายการหนี้ที่ถูกต้องในไฟล์รายงาน",
                "total_records": 0
            }

        # จัดกลุ่มตาม (เลขที่บัญชี, สาขา)
        groups = defaultdict(lambda: {
            "acc": "",
            "branch": "",
            "name_report": "",
            "status": "",
            "periods": [],
            "est_debt": 0.0,
            "real_debt": 0.0,
            "surcharge": 0.0,
            "total_debt": 0.0,
            "steps": [],
            "step_date": "",
            "biz_type": ""
        })

        for rec in raw_records:
            key = f"{rec['acc']};{rec['branch']}"
            g = groups[key]
            g["acc"] = rec["acc"]
            g["branch"] = rec["branch"]
            if not g["name_report"] and rec["name_report"]:
                g["name_report"] = rec["name_report"]
            if not g["status"] and rec["status"]:
                g["status"] = rec["status"]
            if rec["period"]:
                g["periods"].append(rec["period"])
            g["est_debt"] += rec["est_debt"]
            g["real_debt"] += rec["real_debt"]
            g["surcharge"] += rec["surcharge"]
            g["total_debt"] += rec["total_debt"]
            if rec.get("step") and rec["step"] not in g["steps"]:
                g["steps"].append(rec["step"])
            if rec.get("step_date"):
                g["step_date"] = rec["step_date"]
            if not g["biz_type"] and rec["biz_type"]:
                g["biz_type"] = rec["biz_type"]

        # สร้างตารางผลลัพธ์
        processed_list = []
        matched_wcf_count = 0

        for key, g in groups.items():
            wcf_data = self.wcf_integrator.lookup(g["acc"], g["branch"])
            is_matched = bool(wcf_data.get("ชื่อสปก_จริง"))
            if is_matched:
                matched_wcf_count += 1

            # งวดเดือนค้างชำระ (แบบ B)
            period_text = format_period_ranges_b(g["periods"])
            principal_debt = round(g["est_debt"] + g["real_debt"], 2)

            item = {
                "เลขอ้างอิง": key,
                "เลขที่บัญชี": g["acc"],
                "สาขา": g["branch"],
                "ชื่อสถานประกอบการ (จริง)": wcf_data.get("ชื่อสปก_จริง", ""),
                "ที่อยู่": wcf_data.get("ที่อยู่", ""),
                "ตำบล": wcf_data.get("ตำบล", ""),
                "อำเภอ": wcf_data.get("อำเภอ", ""),
                "จังหวัด": wcf_data.get("จังหวัด", ""),
                "รหัสไปรษณีย์": wcf_data.get("รหัสไปรษณีย์", ""),
                "จำนวนงวดค้าง": len(set(g["periods"])),
                "งวดเดือนค้างชำระ": period_text,
                "หนี้ประเมิน": round(g["est_debt"], 2),
                "หนี้จริง": round(g["real_debt"], 2),
                "ยอดหนี้ไม่รวมเงินเพิ่ม": principal_debt,
                "เงินเพิ่ม": round(g["surcharge"], 2),
                "ยอดค้างชำระรวมทั้งสิ้น": round(g["total_debt"], 2),
                "สถานะ สปก.": g["status"],
                "ขั้นตอน": ", ".join(g["steps"]),
                "วันที่บันทึกขั้นตอน": g["step_date"],
                "ประเภทกิจการ": g["biz_type"],
                "ชื่อในรายงานหนี้": g["name_report"],
                "โทรศัพท์": wcf_data.get("โทรศัพท์", ""),
                "เลขพาณิชย์": wcf_data.get("เลขพาณิชย์", "")
            }
            processed_list.append(item)

        df_result = pd.DataFrame(processed_list)

        # จัดเตรียมบันทึกไฟล์ส่งออก
        if not output_dir:
            output_dir = os.path.join(os.path.dirname(__file__), "Output")
        os.makedirs(output_dir, exist_ok=True)

        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        if not custom_filename:
            base_input_name = os.path.splitext(os.path.basename(file_path))[0]
            excel_filename = f"Debt_Processed_{base_input_name}_{timestamp}.xlsx"
        else:
            excel_filename = custom_filename if custom_filename.endswith(".xlsx") else f"{custom_filename}.xlsx"

        output_path = os.path.join(output_dir, excel_filename)
        self.export_to_excel(df_result, output_path)

        # สถิติสรุปผล
        total_principal = float(df_result["ยอดหนี้ไม่รวมเงินเพิ่ม"].sum())
        total_surcharge = float(df_result["เงินเพิ่ม"].sum())
        grand_total = float(df_result["ยอดค้างชำระรวมทั้งสิ้น"].sum())

        return {
            "success": True,
            "total_raw_records": len(raw_records),
            "total_establishments": len(df_result),
            "matched_wcf_count": matched_wcf_count,
            "unmatched_wcf_count": len(df_result) - matched_wcf_count,
            "match_rate_pct": round((matched_wcf_count / len(df_result)) * 100, 1) if len(df_result) > 0 else 0,
            "total_principal_debt": total_principal,
            "total_surcharge": total_surcharge,
            "grand_total_debt": grand_total,
            "output_file": output_path,
            "output_filename": excel_filename,
            "preview_data": df_result.to_dict(orient="records")
        }

    def export_to_excel(self, df: pd.DataFrame, output_path: str):
        """
        จัดรูปแบบและส่งออก DataFrame สู่ Excel อย่างสวยงามและเป็นมืออาชีพ:
        - หัวตารางพื้นหลังสีน้ำเงินเข้ม (SSO Slate Navy) ตัวอักษรสีขาว หนา กึ่งกลาง
        - ตัวเลขจัดชิดขวา พร้อมคอมม่าและทศนิยม 2 ตำแหน่ง
        - รหัสและตัวอักษรจัดชิดซ้ายหรือกึ่งกลางตามความเหมาะสม
        - ปรับความกว้างคอลัมน์อัตโนมัติ ไม่ให้ตัวเลขขึ้น ###
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "สรุปยอดหนี้ค้างชำระ"

        # สไตล์
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Deep Blue
        header_font = Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Tahoma", size=10)
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0")
        )

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        # เขียน Header
        columns = list(df.columns)
        ws.append(columns)
        for col_num in range(1, len(columns) + 1):
            cell = ws.cell(1, col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center
            cell.border = thin_border
        ws.row_dimensions[1].height = 28

        # ระบุคอลัมน์ที่เป็นตัวเงิน
        currency_cols = {"หนี้ประเมิน", "หนี้จริง", "ยอดหนี้ไม่รวมเงินเพิ่ม", "เงินเพิ่ม", "ยอดค้างชำระรวมทั้งสิ้น"}
        integer_cols = {"จำนวนงวดค้าง"}
        center_cols = {"เลขอ้างอิง", "เลขที่บัญชี", "สาขา", "รหัสไปรษณีย์", "สถานะ สปก.", "ขั้นตอน", "วันที่บันทึกขั้นตอน"}

        # เขียนข้อมูลแถว
        for r_idx, row_data in enumerate(df.values, start=2):
            ws.append(list(row_data))
            ws.row_dimensions[r_idx].height = 20
            for c_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(r_idx, c_idx)
                col_name = columns[c_idx - 1]
                cell.font = regular_font
                cell.border = thin_border

                if col_name in currency_cols:
                    cell.number_format = "#,##0.00"
                    cell.alignment = align_right
                elif col_name in integer_cols:
                    cell.number_format = "#,##0"
                    cell.alignment = align_center
                elif col_name in center_cols:
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

        # เขียนแถวสรุปผลรวมท้ายตาราง (Summary Row)
        if len(df) > 0:
            sum_row_idx = len(df) + 2
            ws.row_dimensions[sum_row_idx].height = 24
            summary_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
            summary_font = Font(name="Tahoma", size=10, bold=True, color="0F172A")
            double_bottom_border = Border(
                left=Side(style="thin", color="CBD5E1"),
                right=Side(style="thin", color="CBD5E1"),
                top=Side(style="thin", color="64748B"),
                bottom=Side(style="double", color="0F172A")
            )

            for c_idx, col_name in enumerate(columns, start=1):
                cell = ws.cell(sum_row_idx, c_idx)
                cell.fill = summary_fill
                cell.font = summary_font
                cell.border = double_bottom_border

                if c_idx == 1:
                    cell.value = "รวมทั้งสิ้น"
                    cell.alignment = align_center
                elif col_name in currency_cols:
                    cell.value = round(float(df[col_name].sum()), 2)
                    cell.number_format = "#,##0.00"
                    cell.alignment = align_right
                elif col_name in integer_cols:
                    cell.value = int(df[col_name].sum())
                    cell.number_format = "#,##0"
                    cell.alignment = align_center
                else:
                    cell.value = ""

        # ปรับความกว้างคอลัมน์อัตโนมัติ
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                # ภาษาไทยใช้ความกว้างมากกว่า ASCII เล็กน้อย
                length = sum(2 if ord(char) > 127 else 1 for char in val)
                if length > max_len:
                    max_len = length
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Freeze Panes หัวแถว
        ws.freeze_panes = "A2"

        wb.save(output_path)
