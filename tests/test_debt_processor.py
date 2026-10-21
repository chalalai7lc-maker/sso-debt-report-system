import os
import sys
import unittest
import tempfile
import openpyxl

# Add parent directory to path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from debt_processor import DebtReportProcessor, format_period_ranges_b, parse_float_safe
from wcf_integrator import WCFIntegrator

class TestDebtProcessor(unittest.TestCase):

    def test_parse_float_safe_with_commas(self):
        self.assertEqual(parse_float_safe("220,668.00"), 220668.0)
        self.assertEqual(parse_float_safe("1,234,567.89"), 1234567.89)
        self.assertEqual(parse_float_safe(""), 0.0)
        self.assertEqual(parse_float_safe(None), 0.0)
        self.assertEqual(parse_float_safe(100), 100.0)

    def test_format_period_ranges_b_single(self):
        self.assertEqual(format_period_ranges_b(["07/2569"]), "กรกฎาคม 2569")

    def test_format_period_ranges_b_consecutive(self):
        self.assertEqual(
            format_period_ranges_b(["06/2569", "07/2569"]),
            "มิถุนายน 2569 - กรกฎาคม 2569"
        )

    def test_format_period_ranges_b_non_consecutive(self):
        # 04, 05, 07
        self.assertEqual(
            format_period_ranges_b(["04/2569", "05/2569", "07/2569"]),
            "เมษายน 2569 - พฤษภาคม 2569, กรกฎาคม 2569"
        )

    def test_format_period_ranges_b_cross_year(self):
        self.assertEqual(
            format_period_ranges_b(["12/2568", "01/2569", "02/2569"]),
            "ธันวาคม 2568 - กุมภาพันธ์ 2569"
        )

    def test_format_period_ranges_b_complex(self):
        self.assertEqual(
            format_period_ranges_b(["08/2553", "09/2553", "04/2554", "04/2555"]),
            "สิงหาคม 2553 - กันยายน 2553, เมษายน 2554, เมษายน 2555"
        )

    def test_process_actual_file(self):
        base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
        input_path = os.path.join(base_dir, "Input", "CN200090-1_20260914_161817.xlsx")
        self.assertTrue(os.path.exists(input_path), f"Input file should exist at {input_path}")

        processor = DebtReportProcessor()
        with tempfile.TemporaryDirectory() as tmp_dir:
            result = processor.process(input_path, output_dir=tmp_dir)

            self.assertTrue(result["success"])
            self.assertEqual(result["total_raw_records"], 314)
            self.assertEqual(result["total_establishments"], 152)
            self.assertEqual(result["matched_wcf_count"], 148)
            self.assertEqual(result["unmatched_wcf_count"], 4)
            self.assertAlmostEqual(result["total_principal_debt"], 1885246.00, places=2)
            self.assertAlmostEqual(result["total_surcharge"], 81109.05, places=2)
            self.assertAlmostEqual(result["grand_total_debt"], 1966355.05, places=2)

            # Check output file exists in temp_dir
            output_file = result["output_file"]
            self.assertTrue(os.path.exists(output_file))

            # Check excel output content
            wb = openpyxl.load_workbook(output_file)
            ws = wb.active
            self.assertEqual(ws.title, "สรุปยอดหนี้ค้างชำระ")
            # Header + 152 data rows + 1 summary row = 154 rows
            self.assertEqual(ws.max_row, 154)
            self.assertEqual(ws.cell(154, 1).value, "รวมทั้งสิ้น")

    def test_wcf_auto_reload_on_change(self):
        integrator = WCFIntegrator()
        self.assertTrue(integrator.is_loaded)
        initial_mtime = integrator._last_mtime
        self.assertGreater(initial_mtime, 0)

        # Force mtime check with simulated timestamp
        integrator._last_mtime = initial_mtime - 100
        # When lookup is called, it should detect mtime mismatch and reload
        res = integrator.lookup("4700020784", "000000")
        self.assertEqual(integrator._last_mtime, initial_mtime)
        self.assertIn("อุดรธานี", res["จังหวัด"])

if __name__ == "__main__":
    unittest.main()
