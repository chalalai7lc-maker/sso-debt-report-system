import os
import sys
import unittest
from unittest.mock import patch, AsyncMock
import openpyxl
from fastapi.testclient import TestClient

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import app as app_module
from app import app, OUTPUT_DIR


class TestSSOCoreAndPeriodEdit(unittest.TestCase):

    def setUp(self):
        self.client = TestClient(app)

    def test_process_includes_original_periods_and_source_metadata(self):
        res = self.client.post(
            "/api/process-existing",
            json={"filename": "CN200090-1_20260914_161817.xlsx"}
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertTrue(data["success"])
        first_row = data["preview_data"][0]
        self.assertIn("งวดดิบ", first_row)
        self.assertIn("งวดดิบ_ต้นฉบับ", first_row)
        self.assertEqual(first_row["งวดดิบ"], first_row["งวดดิบ_ต้นฉบับ"])
        self.assertEqual(first_row.get("แหล่งข้อมูลงวด"), "Excel")

        out_f = os.path.join(OUTPUT_DIR, data["output_filename"])
        if os.path.exists(out_f):
            os.remove(out_f)

    def test_sso_core_check_endpoint_updates_overdue_and_paid_accounts(self):
        fake_results = {
            ("4700012345", "000000"): {
                "clean_account_no": "4700012345",
                "clean_branch_code": "000000",
                "company_name": "บริษัท ทดสอบค้างชำระ จำกัด",
                "company_status": "A - เป็นสมาชิก",
                "latest_paid_period": "05/2569",
                "overdue_count": 2,
                "overdue_periods": ["06/2569", "07/2569"],
                "overdue_periods_str": "06/2569, 07/2569",
                "is_suspect": False,
                "reg_timeline_summary": "-",
                "process_status": "ตรวจสอบสำเร็จ"
            },
            ("4700099999", "000000"): {
                "clean_account_no": "4700099999",
                "clean_branch_code": "000000",
                "company_name": "หจก. จ่ายครบแล้ว",
                "company_status": "A - เป็นสมาชิก",
                "latest_paid_period": "09/2569",
                "overdue_count": 0,
                "overdue_periods": [],
                "overdue_periods_str": "-",
                "is_suspect": False,
                "reg_timeline_summary": "-",
                "process_status": "ตรวจสอบสำเร็จ"
            }
        }

        async def mock_check(company_no, branch_code="000000", start_year=2568, as_of_date=None):
            acc = str(company_no).strip().zfill(10)
            br = str(branch_code).strip().zfill(6)
            return fake_results[(acc, br)]

        with patch.object(app_module.sso_client, "get_token", new=AsyncMock(return_value="mock-token")), \
             patch.object(app_module.sso_client, "check_account_overdue", side_effect=mock_check):
            res = self.client.post(
                "/api/sso-core/check",
                json={
                    "start_year": 2568,
                    "accounts": [
                        {"acc": "4700012345", "branch": "000000"},
                        {"acc": "4700099999", "branch": "000000"},
                    ]
                }
            )
            self.assertEqual(res.status_code, 200)
            body = res.json()
            self.assertTrue(body["success"])
            items = body["results"]
            self.assertEqual(len(items), 2)

            r1 = items[0]
            self.assertEqual(r1["acc"], "4700012345")
            self.assertEqual(r1["overdue_periods"], ["06/2569", "07/2569"])
            self.assertEqual(r1["overdue_count"], 2)
            self.assertEqual(r1["period_text_b"], "มิถุนายน 2569 - กรกฎาคม 2569")
            self.assertEqual(r1["sso_check_badge"], "ค้างชำระ 2 งวด")

            r2 = items[1]
            self.assertEqual(r2["acc"], "4700099999")
            self.assertEqual(r2["overdue_periods"], [])
            self.assertEqual(r2["overdue_count"], 0)
            self.assertEqual(r2["period_text_b"], "ชำระครบแล้ว (0 งวด)")
            self.assertEqual(r2["sso_check_badge"], "ชำระครบแล้ว (0 งวด)")

    def test_re_export_endpoint_reflects_edited_periods_in_excel(self):
        sample_records = [
            {
                "เลขอ้างอิง": "4700012345;000000",
                "เลขที่บัญชี": "4700012345",
                "สาขา": "000000",
                "ชื่อสถานประกอบการ (จริง)": "บริษัท ทดสอบแก้ไขงวด จำกัด",
                "ที่อยู่": "123 ถ.นิตโย",
                "ตำบล": "สว่างแดนดิน",
                "อำเภอ": "สว่างแดนดิน",
                "จังหวัด": "สกลนคร",
                "รหัสไปรษณีย์": "47110",
                "จำนวนงวดค้าง": 1,
                "งวดเดือนค้างชำระ": "มกราคม 2569",
                "หนี้ประเมิน": 0.0,
                "หนี้จริง": 1500.0,
                "ยอดหนี้ไม่รวมเงินเพิ่ม": 1500.0,
                "เงินเพิ่ม": 100.0,
                "ยอดค้างชำระรวมทั้งสิ้น": 1600.0,
                "สถานะ สปก.": "A",
                "ขั้นตอน": "เตือน",
                "วันที่บันทึกขั้นตอน": "01/09/2569",
                "ประเภทกิจการ": "11",
                "ชื่อในรายงานหนี้": "บริษัท ทดสอบแก้ไขงวด จำกัด",
                "โทรศัพท์": "042776833",
                "เลขพาณิชย์": "0475560000111",
                "งวดดิบ": ["03/2569", "01/2569", "02/2569"],
                "แหล่งข้อมูลงวด": "แก้ไขเอง",
            }
        ]
        res = self.client.post(
            "/api/records/re-export",
            json={"records": sample_records, "source_filename": "test_edit"}
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertTrue(data["success"])
        self.assertEqual(data["preview_data"][0]["จำนวนงวดค้าง"], 3)
        self.assertEqual(data["preview_data"][0]["งวดเดือนค้างชำระ"], "มกราคม 2569 - มีนาคม 2569")
        self.assertEqual(data["preview_data"][0]["งวดดิบ"], ["01/2569", "02/2569", "03/2569"])

        out_path = os.path.join(OUTPUT_DIR, data["output_filename"])
        self.assertTrue(os.path.exists(out_path))
        wb = openpyxl.load_workbook(out_path, data_only=True)
        ws = wb.active
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        row2 = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        row_dict = dict(zip(headers, row2))
        self.assertEqual(row_dict["จำนวนงวดค้าง"], 3)
        self.assertEqual(row_dict["งวดเดือนค้างชำระ"], "มกราคม 2569 - มีนาคม 2569")
        wb.close()
        os.remove(out_path)

    def test_index_html_has_sso_core_and_edit_period_controls(self):
        res = self.client.get("/")
        self.assertEqual(res.status_code, 200)
        html = res.text
        self.assertIn("id=\"sso-check-btn\"", html)
        self.assertIn("id=\"sso-start-year\"", html)
        self.assertIn("id=\"edit-periods-modal\"", html)


if __name__ == "__main__":
    unittest.main()

